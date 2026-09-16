from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

def load_pdf_file(file_path: str) -> str:
    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text

def load_text_file(file_path: str) -> str:
    path = Path(file_path)

    with open(path, "r", encoding="utf-8") as file:
        text = file.read()

    return text

if __name__ == "__main__":
    text = load_text_file("data/company.txt")

    print(text)

if __name__ == "__main__":
    text = load_pdf_file("data/Geometric Distribution.pdf")

    print("\n===== PDF TEXT =====")
    print(text[:3000])

def load_docx_file(file_path: str) -> str:
    document = Document(file_path)

    text = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    return "\n".join(text)

if __name__ == "__main__":
    text = load_docx_file("data/Climate Change.docx")

    print("\n===== DOCX TEXT =====")
    print(text[:3000])

def load_excel_file(file_path: str) -> str:
    workbook = load_workbook(
        file_path,
        data_only=True
    )

    text = []

    for sheet in workbook.worksheets:

        text.append(f"Sheet: {sheet.title}")

        for row in sheet.iter_rows(values_only=True):

            values = [
                str(value)
                for value in row
                if value is not None
            ]

            if values:
                text.append(" | ".join(values))

    return "\n".join(text)

if __name__ == "__main__":
    text = load_excel_file("data/02 Project Management.xlsx")

    print("\n===== EXCEL TEXT =====")
    print(text[:3000])

def load_document(file_path: str) -> str:
    path = Path(file_path)

    extension = path.suffix.lower()

    if extension == ".txt":
        return load_text_file(file_path)

    elif extension == ".pdf":
        return load_pdf_file(file_path)

    elif extension == ".docx":
        return load_docx_file(file_path)

    elif extension == ".xlsx":
        return load_excel_file(file_path)

    else:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )